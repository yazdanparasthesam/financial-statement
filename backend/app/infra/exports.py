"""XLSX (openpyxl) and PDF (fpdf2) exports of generated statements."""
import io

_SAFE = {"\u2014": "-", "\u2013": "-", "\u2018": "'", "\u2019": "'",
         "\u201c": '"', "\u201d": '"', "\u2026": "...", "\u00a0": " "}


def _safe(text: str) -> str:
    """fpdf2 core fonts are latin-1: replace typographic chars, strip the rest."""
    for k, v in _SAFE.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "ignore").decode("latin-1")


def _flatten(kind: str, data: dict) -> list[tuple[str, float]]:
    rows: list[tuple[str, float]] = []
    title = "Profit & Loss" if kind == "pl" else "Balance Sheet"
    rows.append((title, 0.0))
    for key, node in data.items():
        if isinstance(node, dict) and "items" in node:
            rows.append((key.replace("_", " ").title(), node["total"]))
            for item in node["items"]:
                rows.append((f"    {item['code']} {item['name']}", item["amount"]))
        elif isinstance(node, (int, float)):
            rows.append((key.replace("_", " ").title(), float(node)))
    return rows


def statement_xlsx(kind: str, data: dict, meta: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "PL" if kind == "pl" else "BS"
    ws.append(["Financial Statement — statutory statement export"])
    ws.append([f"Company: {meta.get('company', '')}", f"Fiscal year: {meta.get('fiscal_year', '')}",
               f"Version: {meta.get('version', '')}"])
    ws.append([])
    ws.append(["Line", "Amount"])
    for label, amount in _flatten(kind, data):
        ws.append([label, amount])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def statement_pdf(kind: str, data: dict, meta: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, _safe("Financial Statement - statutory statement export"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, _safe(f"Company: {meta.get('company', '')}   Fiscal year: {meta.get('fiscal_year', '')}"
                          f"   Version: {meta.get('version', '')}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for label, amount in _flatten(kind, data):
        bold = not label.startswith("    ")
        pdf.set_font("Helvetica", "B" if bold else "", 10)
        pdf.cell(130, 6, _safe(label.strip() if not bold else label), border=0)
        pdf.cell(50, 6, f"{amount:,.2f}" if (bold or amount) else "", align="R",
                 new_x="LMARGIN", new_y="NEXT")
    out = pdf.output()
    return bytes(out)


def checklist_pdf(items: list[dict], meta: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Filing checklist", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, _safe(f"Company: {meta.get('company', '')}   Fiscal year: {meta.get('fiscal_year', '')}"),
             new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for it in items:
        pdf.cell(8, 6, "X" if it["done"] else " ")
        pdf.cell(0, 6, _safe(it["item"]), new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
