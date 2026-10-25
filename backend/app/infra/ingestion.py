"""Trial-balance ingestion: CSV and XLSX → normalized rows. Friendly errors, no pandas."""
import csv
import io


class IngestionError(ValueError):
    pass


CODE_KEYS = ("code", "account code", "account_code", "acc code", "konto")
NAME_KEYS = ("name", "description", "account name", "account", "title", "hesap")
DEBIT_KEYS = ("debit", "deb", "bedehkar", "dr")
CREDIT_KEYS = ("credit", "cred", "cre", "bestankar", "cr")


def _find(header: list[str], keys: tuple[str, ...]) -> int | None:
    norm = [h.strip().lower() for h in header]
    for k in keys:
        if k in norm:
            return norm.index(k)
    return None


def _to_float(value) -> float:
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(",", "")
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        raise IngestionError(f"Cannot parse amount: {value!r}") from None


def _rows_from_csv(content: bytes) -> list[list]:
    text = content.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]


def _rows_from_xlsx(content: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise IngestionError("openpyxl is required to parse XLSX files") from e
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [[("" if c is None else c) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return [r for r in rows if any(str(c).strip() for c in r)]


def parse_trial_balance(filename: str, content: bytes) -> list[dict]:
    """Return [{code,name,debit,credit}, ...]. Raises IngestionError with guidance."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(content)
    elif lower.endswith((".csv", ".txt")):
        rows = _rows_from_csv(content)
    else:
        raise IngestionError("Unsupported file type. Upload a .csv or .xlsx trial balance.")
    if len(rows) < 2:
        raise IngestionError("File has no data rows. Expected a header plus account rows.")

    header, data = rows[0], rows[1:]
    ci, ni = _find(header, CODE_KEYS), _find(header, NAME_KEYS)
    di, cri = _find(header, DEBIT_KEYS), _find(header, CREDIT_KEYS)
    missing = [n for n, i in (("code", ci), ("name", ni), ("debit", di), ("credit", cri)) if i is None]
    if missing:
        raise IngestionError(
            f"Could not find column(s): {', '.join(missing)}. "
            f"Header must include code, name, debit, credit (got: {header})."
        )

    accounts = []
    for row in data:
        code = str(row[ci]).strip()
        if not code:
            continue
        accounts.append({
            "code": code,
            "name": str(row[ni]).strip(),
            "debit": _to_float(row[di] if di < len(row) else 0),
            "credit": _to_float(row[cri] if cri < len(row) else 0),
        })
    if not accounts:
        raise IngestionError("No account rows with a code were found.")
    return accounts
